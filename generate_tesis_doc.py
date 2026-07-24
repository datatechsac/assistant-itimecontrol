from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "logs" / "rag_evaluation_results.json"
OUTPUT = ROOT / "tesis_apartados_rag.docx"

with INPUT.open("r", encoding="utf-8") as f:
    data = json.load(f)

summary = data.get("summary", {})

# Build document
document = Document()


def add_image_with_caption(doc, caption: str, image_path: Path, width=Inches(6.0)) -> None:
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(caption)
    run.bold = True
    doc.add_paragraph()
    doc.add_picture(str(image_path), width=width)

# Style
style = document.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# Title
p = document.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Propuesta de solución basada en RAG para iTimeControl")
run.bold = True
run.font.size = Pt(14)

p = document.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Apartados para tesis: Desarrollo, resultados, discusión e impacto")
run.italic = True
run.font.size = Pt(11)

# Section 3.2
heading = document.add_heading("3.2 Desarrollo de la Propuesta de Solución", level=1)
paragraph = document.add_paragraph()
paragraph.add_run(
    "El desarrollo de la propuesta se estructuró en cuatro etapas complementarias: recolección y preparación de datos, construcción del índice vectorial, diseño del pipeline de recuperación y generación, y evaluación del sistema. Cada una de estas etapas permitió transformar un conjunto de documentos del dominio de iTimeControl en un asistente conversacional capaz de responder preguntas en lenguaje natural con base documental."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "En la primera etapa, se recopilaron y organizaron documentos relevantes del sistema iTimeControl, incluyendo manuales de usuario, preguntas frecuentes, documentación técnica y material normativo. Posteriormente, se realizó una limpieza de texto para eliminar ruido, encabezados redundantes y elementos no pertinentes, con el objetivo de preservar únicamente información útil para el modelo. Luego, los documentos fueron segmentados en chunks con superposición, lo que permitió conservar contexto semántico y mejorar la recuperación de información relevante."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "En la segunda etapa, se construyó el componente de recuperación semántica. Para ello, se generaron embeddings a partir de los chunks mediante un modelo multilingual de sentence-transformers y se almacenaron en una base vectorial FAISS. Esta etapa permitió representar el contenido documental en un espacio vectorial, de manera que una consulta del usuario pudiera compararse con los documentos almacenados y recuperar los fragmentos más similares."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "En la tercera etapa, se diseñó el pipeline RAG. El flujo consistió en recibir una pregunta del usuario, recuperar los chunks más relevantes, construir un prompt enriquecido con el contexto recuperado y enviarlo a un modelo de lenguaje para generar una respuesta grounded en la información documental. Esta arquitectura se implementó de forma modular mediante módulos de preprocessing, rag, evaluation y api, lo que facilitó la reproducibilidad, el mantenimiento y la futura escalabilidad del sistema."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "En la cuarta etapa, se llevó a cabo la evaluación del sistema con un conjunto de 20 preguntas de benchmark, seleccionadas para representar escenarios reales de consulta sobre iTimeControl. Los resultados obtenidos mostraron que el sistema fue capaz de responder de forma contextualizada y útil, aunque con un margen de mejora en la precisión de recuperación. En términos cuantitativos, el modelo alcanzó un ROUGE-1 de 0.2835, un ROUGE-2 de 0.0846, un ROUGE-L de 0.1849 y un BLEU de 0.0707. Asimismo, la métrica de Exact Match fue de 0.0000, lo que indica que no hubo coincidencia literal exacta con las respuestas de referencia. En cuanto a recuperación, el Hit Rate fue de 0.3500, el Context Recall alcanzó 0.4578, el MRR fue 0.2000 y los valores de Recall@1, Recall@3, Recall@5 y Recall@10 fueron 0.1000, 0.2500, 0.3500 y 0.3500, respectivamente."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "Estos resultados sugieren que la propuesta funciona como una base sólida para el desarrollo de un asistente documental especializado, aunque aún requiere mejoras en la calidad de la recuperación y en la capacidad de seleccionar de forma más precisa los fragmentos más pertinentes para cada consulta."
)

# Section 3.3
heading = document.add_heading("3.3 Análisis de los Datos y Resultados", level=1)
paragraph = document.add_paragraph()
paragraph.add_run(
    "Para evaluar la propuesta se utilizó un conjunto de preguntas de benchmark compuesto por 20 consultas orientadas a temas de iTimeControl. La evaluación consideró métricas de generación textual y métricas de recuperación, con el fin de medir no solo la calidad del texto generado, sino también la capacidad del sistema para localizar información relevante dentro del corpus documental."
)

# Table
table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.autofit = True
cells = table.rows[0].cells
cells[0].text = "Métrica"
cells[1].text = "Valor"
cells[2].text = "Observación"

rows = [
    ("ROUGE-1", f"{summary.get('rouge1', 0):.4f}", "Nivel moderado de similitud léxica con la referencia"),
    ("ROUGE-2", f"{summary.get('rouge2', 0):.4f}", "Baja coincidencia de bigramas, lo que indica reformulación del contenido"),
    ("ROUGE-L", f"{summary.get('rougeL', 0):.4f}", "Cobertura parcial de la estructura semántica de la respuesta esperada"),
    ("BLEU", f"{summary.get('bleu', 0):.4f}", "Puntaje bajo, lo que refleja respuestas menos literales que las referencias"),
    ("Exact Match", f"{summary.get('exact_match', 0):.4f}", "Sin coincidencia exacta en el conjunto evaluado"),
    ("Hit Rate", f"{summary.get('hit_rate', 0):.4f}", "El 35% de las preguntas recuperó al menos un chunk relevante"),
    ("Context Recall", f"{summary.get('context_recall', 0):.4f}", "El sistema recuperó de forma parcial el contexto esperado"),
    ("MRR", f"{summary.get('mrr', 0):.4f}", "La información relevante aparece en posiciones medias del ranking"),
    ("Recall@1", f"{summary.get('recall@1', 0):.4f}", "Recuperación precisa en el primer resultado en el 10% de los casos"),
    ("Recall@3", f"{summary.get('recall@3', 0):.4f}", "El 25% de las consultas alcanzó un hit dentro de los tres primeros resultados"),
    ("Recall@5", f"{summary.get('recall@5', 0):.4f}", "El 35% de las consultas tuvo una recuperación útil dentro de los cinco primeros resultados"),
    ("Recall@10", f"{summary.get('recall@10', 0):.4f}", "Se mantiene el mismo nivel de cobertura en rangos más amplios de recuperación"),
]

for metric, value, note in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = metric
    row_cells[1].text = value
    row_cells[2].text = note

paragraph = document.add_paragraph()
paragraph.add_run(
    "Los resultados indican que el modelo puede responder de manera útil y contextualizada, aunque todavía presenta margen de mejora en la precisión de recuperación. La métrica de Context Recall, cercana al 46%, sugiere que el sistema recupera parcialmente la información relevante, mientras que el Hit Rate del 35% evidencia que la respuesta correcta o suficientemente alineada no siempre aparece dentro del conjunto de documentos seleccionados."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "Asimismo, los resultados sugieren que las respuestas generadas son coherentes y pertinentes en muchos casos, pero no siempre literalmente coinciden con las referencias esperadas. Esto se observa en los valores bajos de ROUGE-2, BLEU y Exact Match, lo cual es esperable cuando el modelo reformula información en lenguaje natural en vez de copiar literalmente el texto fuente."
)

add_image_with_caption(
    document,
    "Figura 1. Rendimiento general del sistema RAG",
    ROOT / "logs" / "rag_metrics_bar.png",
)
add_image_with_caption(
    document,
    "Figura 2. Recall@k obtenido por el sistema de recuperación",
    ROOT / "logs" / "rag_recall_at_k.png",
)
add_image_with_caption(
    document,
    "Figura 3. Matriz de métricas del modelo RAG",
    ROOT / "logs" / "rag_metrics_heatmap.png",
)
add_image_with_caption(
    document,
    "Figura 4. Distribución de fuentes recuperadas por el retriever",
    ROOT / "logs" / "retriever_sources.png",
)

# Section 3.4
heading = document.add_heading("3.4 Discusión e Interpretación de los Resultados", level=1)
paragraph = document.add_paragraph()
paragraph.add_run(
    "Los resultados más relevantes del experimento muestran que la arquitectura RAG propuesta funciona como una base sólida para la construcción de un asistente especializado en iTimeControl. La combinación entre recuperación semántica y generación con lenguaje natural permitió obtener respuestas contextualizadas, lo que representa una ventaja frente a un enfoque puramente generativo sin acceso a la base documental."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "La interpretación de los resultados sugiere que la calidad del sistema depende en gran medida de la calidad de la recuperación de información. Cuando el retriever encuentra los chunks correctos, la respuesta generada tiende a ser más útil y alineada con la pregunta. Sin embargo, cuando la recuperación falla o no incluye el contexto adecuado, el modelo puede ofrecer respuestas vagas o incompletas. Este patrón es consistente con la literatura sobre sistemas RAG, donde la generación solo puede ser tan buena como la información recuperada."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "Además, los resultados revelan que la estrategia de chunking y la selección de documentos son determinantes para el desempeño. El uso de chunks con contexto y la organización del corpus en documentos del dominio permitieron que el modelo respondiera de forma más precisa. No obstante, las métricas muestran que aún existe margen para mejorar la precisión de los documentos recuperados, especialmente en preguntas ambiguas o de alto contexto legal."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "Entre las limitaciones del estudio se encuentra el tamaño reducido del conjunto de evaluación, compuesto por 20 preguntas de benchmark, lo que limita la generalización de los resultados. Asimismo, la evaluación se centró en métricas automáticas y no incluyó una evaluación humana amplia, por lo que no se mide completamente la utilidad percibida por usuarios reales. Otra limitación relevante es la dependencia de un proveedor externo de generación de respuestas, lo que puede afectar la estabilidad y reproducibilidad del sistema en entornos cambiantes."
)

# Section 3.5
heading = document.add_heading("3.5 Estimación del Impacto de la Solución", level=1)
paragraph = document.add_paragraph()
paragraph.add_run(
    "La solución propuesta tiene un impacto potencial significativo en la operatividad de las organizaciones que utilizan iTimeControl, especialmente en tareas relacionadas con consulta de información, capacitación de usuarios y apoyo a la toma de decisiones operativas. El asistente permite reducir el tiempo de búsqueda de respuestas en manuales y documentos dispersos, automatizar la consulta frecuente sobre procedimientos y disminuir la dependencia de personas con conocimiento especializado para responder dudas rutinarias."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "Desde una perspectiva organizacional, el impacto esperado se concentran en tres dimensiones principales: eficiencia operativa, reducción de errores de información y mejora en la accesibilidad del conocimiento. Al centralizar la información en un sistema conversacional, los usuarios pueden consultar preguntas en lenguaje natural sin necesidad de revisar extensos manuales o contactar directamente a un experto. Esto contribuye a una mayor agilidad en la resolución de dudas y a una mejor experiencia de uso del sistema."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "También es posible estimar un impacto técnico relevante, ya que la arquitectura propuesta establece una base escalable para futuras mejoras, como la integración con otros módulos de RR. HH., la incorporación de modelos más robustos, el uso de evaluaciones humanas y la expansión del corpus documental. En este sentido, el proyecto no solo entrega una solución funcional, sino también un marco reusable para desarrollar asistentes inteligentes orientados a procesos empresariales específicos."
)

paragraph = document.add_paragraph()
paragraph.add_run(
    "En términos generales, el impacto de la solución puede considerarse preliminarmente alto en cuanto a valor práctico, siempre que se continúe mejorando la precisión de recuperación y se amplíe la evaluación con datos reales de uso. El sistema representa una primera aproximación sólida hacia un asistente documental inteligente, con potencial para convertirse en un componente estratégico dentro del ecosistema de iTimeControl."
)

# Closing
paragraph = document.add_paragraph()
paragraph.add_run("Documento generado automáticamente a partir de los resultados del proyecto y del archivo de evaluación RAG.")
paragraph.italic = True

# Save
document.save(OUTPUT)
print(f"Documento creado en: {OUTPUT}")
